"""Scan data/skills_curation/ for xlsx/csv and upsert skill_catalog."""

from __future__ import annotations

import csv
import logging
import re
from pathlib import Path
from typing import Any, Iterable, Optional

from sqlalchemy.orm import Session

from app.models import SkillCatalog

logger = logging.getLogger("crede.skills_import")

API_ROOT = Path(__file__).resolve().parents[2]
CURATION_DIR = API_ROOT / "data" / "skills_curation"

NAME_KEYS = {"skill", "name", "skill name", "title", "competency", "label"}
CAT_KEYS = {"category", "type", "group", "domain"}
DESC_KEYS = {"description", "definition", "summary"}
ALIAS_KEYS = {"aliases", "also known as", "synonyms", "alternate names"}


def _norm(h: Any) -> str:
    return re.sub(r"\s+", " ", str(h or "").strip().lower())


def _norm_name(t: str) -> str:
    return re.sub(r"\s+", " ", (t or "").strip().lower())


def _pick(row: dict, keys: set) -> Optional[str]:
    for k, v in row.items():
        if _norm(k) in keys and v is not None and str(v).strip():
            return str(v).strip()
    return None


def _split(text: Optional[str]) -> list[str]:
    if not text:
        return []
    return [p.strip() for p in re.split(r"[;|,]", text) if p.strip()][:20]


def discover_files(directory: Path | None = None) -> list[Path]:
    d = directory or CURATION_DIR
    if not d.is_dir():
        return []
    return sorted(
        p for p in d.iterdir()
        if p.suffix.lower() in (".xlsx", ".xls", ".csv") and not p.name.startswith(".")
    )


def _iter_file(path: Path) -> Iterable[dict]:
    if path.suffix.lower() in (".xlsx", ".xls"):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        for sheet in wb.worksheets:
            rows = sheet.iter_rows(values_only=True)
            try:
                headers = next(rows)
            except StopIteration:
                continue
            headers = [str(h) if h is not None else f"c{i}" for i, h in enumerate(headers)]
            for row in rows:
                if not row or all(c is None or str(c).strip() == "" for c in row):
                    continue
                yield dict(zip(headers, row))
        wb.close()
    else:
        with path.open(newline="", encoding="utf-8-sig", errors="replace") as f:
            yield from csv.DictReader(f)


def import_all(db: Session, directory: Path | None = None) -> dict:
    files = discover_files(directory)
    results = {}
    total = 0
    for path in files:
        n = 0
        try:
            for raw in _iter_file(path):
                name = _pick(raw, NAME_KEYS)
                if not name:
                    continue
                norm = _norm_name(name)
                row = db.query(SkillCatalog).filter(SkillCatalog.name_normalized == norm).first()
                data = {
                    "name": name[:120],
                    "name_normalized": norm[:120],
                    "category": (_pick(raw, CAT_KEYS) or "technical")[:64],
                    "description": _pick(raw, DESC_KEYS),
                    "aliases": _split(_pick(raw, ALIAS_KEYS)) or None,
                    "source_file": path.name[:255],
                    "is_active": True,
                }
                if row:
                    for k, v in data.items():
                        setattr(row, k, v)
                else:
                    db.add(SkillCatalog(**data))
                n += 1
            db.commit()
            results[path.name] = n
            total += n
        except Exception as e:
            db.rollback()
            results[path.name] = f"error: {e}"
            logger.exception("skills import %s", path)
    return {"files": results, "total_rows": total, "folder": str(directory or CURATION_DIR)}
