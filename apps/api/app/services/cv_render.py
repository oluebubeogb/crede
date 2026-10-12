"""
Render CV / cover letter / bio to real PDF (reportlab) and DOCX (python-docx)
and store under VPS STORAGE_ROOT.
"""

from __future__ import annotations

import io
from pathlib import Path
from typing import Any, Optional

from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT, TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
    HRFlowable,
)

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.services.storage import save_bytes


def _styles():
    base = getSampleStyleSheet()
    styles = {
        "name": ParagraphStyle(
            "Name",
            parent=base["Heading1"],
            fontSize=18,
            textColor=colors.HexColor("#0F172A"),
            spaceAfter=2,
            alignment=TA_LEFT,
        ),
        "headline": ParagraphStyle(
            "Headline",
            parent=base["Normal"],
            fontSize=11,
            textColor=colors.HexColor("#2563EB"),
            spaceAfter=4,
        ),
        "meta": ParagraphStyle(
            "Meta",
            parent=base["Normal"],
            fontSize=9,
            textColor=colors.HexColor("#64748B"),
            spaceAfter=8,
        ),
        "section": ParagraphStyle(
            "Section",
            parent=base["Heading2"],
            fontSize=12,
            textColor=colors.HexColor("#0F172A"),
            spaceBefore=12,
            spaceAfter=4,
        ),
        "body": ParagraphStyle(
            "Body",
            parent=base["Normal"],
            fontSize=9,
            textColor=colors.HexColor("#334155"),
            leading=12,
            spaceAfter=3,
        ),
        "job_title": ParagraphStyle(
            "JobTitle",
            parent=base["Normal"],
            fontSize=10,
            textColor=colors.HexColor("#0F172A"),
            fontName="Helvetica-Bold",
            spaceAfter=1,
        ),
        "job_meta": ParagraphStyle(
            "JobMeta",
            parent=base["Normal"],
            fontSize=8,
            textColor=colors.HexColor("#64748B"),
            spaceAfter=2,
        ),
        "bullet": ParagraphStyle(
            "Bullet",
            parent=base["Normal"],
            fontSize=9,
            textColor=colors.HexColor("#334155"),
            leftIndent=12,
            spaceAfter=1,
        ),
    }
    return styles


def render_pdf(snapshot: dict[str, Any], template: str = "professional") -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=16 * mm,
        bottomMargin=16 * mm,
    )
    styles = _styles()
    story = []

    name = snapshot.get("name") or "Professional"
    story.append(Paragraph(str(name), styles["name"]))
    if snapshot.get("headline"):
        story.append(Paragraph(str(snapshot["headline"]), styles["headline"]))
    meta_parts = [p for p in [snapshot.get("location"), snapshot.get("email")] if p]
    if meta_parts:
        story.append(Paragraph(" · ".join(str(p) for p in meta_parts), styles["meta"]))
    story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#E2E8F0"), spaceAfter=6))

    if snapshot.get("summary"):
        story.append(Paragraph("Summary", styles["section"]))
        story.append(Paragraph(str(snapshot["summary"]), styles["body"]))

    skills = snapshot.get("skills") or []
    if skills:
        story.append(Paragraph("Skills", styles["section"]))
        story.append(Paragraph(", ".join(str(s) for s in skills), styles["body"]))

    experience = snapshot.get("experience") or []
    if experience:
        story.append(Paragraph("Experience", styles["section"]))
        for e in experience:
            title = e.get("job_title") or ""
            org = e.get("organization") or ""
            story.append(Paragraph(f"{title} — {org}", styles["job_title"]))
            dates = f"{e.get('start', '')} – {e.get('end', 'Present')}"
            story.append(Paragraph(dates, styles["job_meta"]))
            if e.get("description"):
                story.append(Paragraph(str(e["description"]), styles["body"]))
            for a in (e.get("achievements") or [])[:6]:
                story.append(Paragraph(f"• {a}", styles["bullet"]))
            for r in (e.get("responsibilities") or [])[:4]:
                story.append(Paragraph(f"• {r}", styles["bullet"]))
            story.append(Spacer(1, 4))

    certs = snapshot.get("certifications") or []
    if certs:
        story.append(Paragraph("Certifications", styles["section"]))
        for c in certs:
            line = c.get("name") or ""
            if c.get("issuer"):
                line += f" — {c['issuer']}"
            story.append(Paragraph(line, styles["body"]))

    doc.build(story)
    return buf.getvalue()


def render_docx(snapshot: dict[str, Any], template: str = "professional") -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    name = snapshot.get("name") or "Professional"
    p = doc.add_paragraph()
    run = p.add_run(str(name))
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    if snapshot.get("headline"):
        p = doc.add_paragraph()
        run = p.add_run(str(snapshot["headline"]))
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)

    meta_parts = [p for p in [snapshot.get("location"), snapshot.get("email")] if p]
    if meta_parts:
        p = doc.add_paragraph()
        run = p.add_run(" · ".join(str(x) for x in meta_parts))
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    def add_heading(text: str):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    if snapshot.get("summary"):
        add_heading("Summary")
        p = doc.add_paragraph(str(snapshot["summary"]))
        for run in p.runs:
            run.font.size = Pt(9)

    skills = snapshot.get("skills") or []
    if skills:
        add_heading("Skills")
        p = doc.add_paragraph(", ".join(str(s) for s in skills))
        for run in p.runs:
            run.font.size = Pt(9)

    experience = snapshot.get("experience") or []
    if experience:
        add_heading("Experience")
        for e in experience:
            p = doc.add_paragraph()
            run = p.add_run(f"{e.get('job_title', '')} — {e.get('organization', '')}")
            run.bold = True
            run.font.size = Pt(10)
            p = doc.add_paragraph()
            run = p.add_run(f"{e.get('start', '')} – {e.get('end', 'Present')}")
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
            if e.get("description"):
                p = doc.add_paragraph(str(e["description"]))
                for run in p.runs:
                    run.font.size = Pt(9)
            for a in (e.get("achievements") or [])[:6]:
                p = doc.add_paragraph(str(a), style="List Bullet")
                for run in p.runs:
                    run.font.size = Pt(9)

    certs = snapshot.get("certifications") or []
    if certs:
        add_heading("Certifications")
        for c in certs:
            line = c.get("name") or ""
            if c.get("issuer"):
                line += f" — {c['issuer']}"
            p = doc.add_paragraph(line)
            for run in p.runs:
                run.font.size = Pt(9)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def write_generated_files(
    profile_id: int,
    doc_id: int,
    snapshot: dict[str, Any],
    template: str = "professional",
) -> tuple[str, str]:
    """Render PDF + DOCX and store on VPS. Returns (pdf_key, docx_key)."""
    pdf_bytes = render_pdf(snapshot, template)
    docx_bytes = render_docx(snapshot, template)
    pdf_key = f"generated/{profile_id}/{doc_id}.pdf"
    docx_key = f"generated/{profile_id}/{doc_id}.docx"
    save_bytes(pdf_key, pdf_bytes)
    save_bytes(docx_key, docx_bytes)
    return pdf_key, docx_key
