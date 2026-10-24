"""Central career import: extract text from an uploaded file, parse it into career facts (AI first,
heuristics as fallback; LinkedIn-style CSV/XLSX exports are read structurally), dedupe against the
profile, and apply the items the user approved. Files are parsed in memory and discarded.
"""
from __future__ import annotations

import csv
import io
import json
import logging
import re
import zipfile
from datetime import date
from typing import Any

from sqlalchemy.orm import Session

from app.models import Certification, Profile, Skill, TimelineEntry
from app.services import llm, vocab

log = logging.getLogger("crede.import")

SUPPORTED = {"pdf", "docx", "doc", "rtf", "txt", "md", "markdown", "csv", "tsv", "json", "xlsx", "odt", "html", "htm"}
MAX_TEXT = 60000

# ─── Text extraction ──────────────────────────────────────────────────────────


def ext_of(name: str) -> str:
    return (name or "").rsplit(".", 1)[-1].lower() if "." in (name or "") else ""


def _strip_rtf(s: str) -> str:
    s = re.sub(r"\\par[d]?\b", "\n", s)
    s = re.sub(r"\\'[0-9a-f]{2}", "", s)
    s = re.sub(r"\\[a-zA-Z]+-?\d* ?", "", s)
    s = re.sub(r"[{}]", "", s)
    return re.sub(r"\n{3,}", "\n\n", s)


def _xml_text(xml: bytes) -> str:
    t = re.sub(rb"</(w:p|text:p|text:h)>", b"\n", xml)
    t = re.sub(rb"<[^>]+>", b"", t)
    return t.decode("utf-8", "ignore")


def extract(data: bytes, ext: str) -> tuple[str, list[dict] | None]:
    """Returns (plain_text, table_rows_or_None). Raises ValueError on unreadable input."""
    if ext == "pdf":
        if not data.startswith(b"%PDF"):
            raise ValueError("That file is not a valid PDF")
        from pypdf import PdfReader
        r = PdfReader(io.BytesIO(data))
        text = "\n".join((p.extract_text() or "") for p in r.pages[:30])
        if not text.strip():
            raise ValueError("This PDF has no selectable text (it may be a scan). Try a DOCX or text export.")
        return text[:MAX_TEXT], None
    if ext == "docx":
        import docx
        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells if c.text.strip()))
        return "\n".join(parts)[:MAX_TEXT], None
    if ext == "odt":
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            return _xml_text(z.read("content.xml"))[:MAX_TEXT], None
    if ext == "doc":  # legacy binary Word: pull readable runs (cheap, best effort)
        runs = re.findall(rb"[\x20-\x7e\r\n\t]{4,}", data)
        text = "\n".join(r.decode("latin-1") for r in runs)
        text = "\n".join(l for l in text.splitlines() if len(re.findall(r"[A-Za-z]", l)) > len(l) * 0.5)
        if len(text) < 80:
            raise ValueError("Couldn't read this .doc file — save it as DOCX or PDF and try again")
        return text[:MAX_TEXT], None
    if ext == "rtf":
        return _strip_rtf(data.decode("latin-1", "ignore"))[:MAX_TEXT], None
    if ext in ("html", "htm"):
        t = re.sub(r"(?is)<(script|style).*?</\1>", "", data.decode("utf-8", "ignore"))
        t = re.sub(r"(?i)<br\s*/?>|</(p|div|li|h\d|tr)>", "\n", t)
        return re.sub(r"<[^>]+>", "", t)[:MAX_TEXT], None
    if ext in ("csv", "tsv"):
        raw = data.decode("utf-8-sig", "ignore")
        # LinkedIn exports sometimes start with a "Notes:" preamble before the header row
        lines = raw.splitlines()
        start = next((i for i, l in enumerate(lines[:8]) if l.count(",") >= 1 and not l.lower().startswith("notes")), 0)
        rows = list(csv.DictReader(io.StringIO("\n".join(lines[start:])), delimiter="\t" if ext == "tsv" else ","))
        return raw[:MAX_TEXT], rows[:500]
    if ext == "xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        rows: list[dict] = []
        text = []
        for ws in wb.worksheets[:6]:
            it = ws.iter_rows(values_only=True)
            header = next(it, None)
            if not header:
                continue
            hdr = [str(h or "").strip() for h in header]
            text.append(f"# {ws.title}\n" + " | ".join(hdr))
            for r in list(it)[:400]:
                d = {hdr[i]: ("" if v is None else str(v)) for i, v in enumerate(r) if i < len(hdr) and hdr[i]}
                if any(d.values()):
                    d["__sheet"] = ws.title
                    rows.append(d)
                    text.append(" | ".join(v for v in d.values() if v))
        return "\n".join(text)[:MAX_TEXT], rows
    if ext == "json":
        obj = json.loads(data.decode("utf-8", "ignore"))
        return json.dumps(obj, ensure_ascii=False, indent=1)[:MAX_TEXT], None
    # txt / md / anything texty
    return data.decode("utf-8", "ignore")[:MAX_TEXT], None


# ─── Date helpers ─────────────────────────────────────────────────────────────

MONTHS = {m: i + 1 for i, m in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"])}
_MON = r"(?:jan(?:uary)?|feb(?:ruary)?|mar(?:ch)?|apr(?:il)?|may|jun(?:e)?|jul(?:y)?|aug(?:ust)?|sep(?:t(?:ember)?)?|oct(?:ober)?|nov(?:ember)?|dec(?:ember)?)"
_DATE = rf"(?:{_MON}\.?\s+\d{{4}}|\d{{1,2}}[/.-]\d{{4}}|\d{{4}}[/.-]\d{{1,2}}|\d{{4}})"
_PRESENT = r"(?:present|current|now|date|till date|to date|ongoing)"
RANGE_RE = re.compile(rf"(?i)\b({_DATE})\s*(?:-|–|—|to|until|~)\s*({_DATE}|{_PRESENT})\b")


def parse_date(s: Any) -> date | None:
    if s is None:
        return None
    if isinstance(s, date):
        return s
    t = str(s).strip().lower().rstrip(".")
    if not t or re.fullmatch(_PRESENT, t):
        return None
    m = re.match(rf"({_MON})\.?\s+(\d{{4}})", t)
    if m:
        return date(int(m.group(2)), MONTHS[m.group(1)[:3]], 1)
    m = re.match(r"(\d{4})[/.-](\d{1,2})(?:[/.-]\d{1,2})?$", t)
    if m and 1 <= int(m.group(2)) <= 12:
        return date(int(m.group(1)), int(m.group(2)), 1)
    m = re.match(r"(\d{1,2})[/.-](\d{4})$", t)
    if m and 1 <= int(m.group(1)) <= 12:
        return date(int(m.group(2)), int(m.group(1)), 1)
    m = re.search(r"\b(19[5-9]\d|20\d{2})\b", t)
    if m:
        return date(int(m.group(1)), 1, 1)
    return None


def _is_present(s: Any) -> bool:
    return bool(s) and bool(re.fullmatch(_PRESENT, str(s).strip().lower()))


# ─── Structured (LinkedIn-style CSV / XLSX) ──────────────────────────────────


def _k(row: dict, *names: str) -> str:
    low = {str(k).strip().lower(): v for k, v in row.items()}
    for n in names:
        v = low.get(n)
        if v not in (None, ""):
            return str(v).strip()
    return ""


def from_rows(rows: list[dict]) -> dict | None:
    out = _empty()
    for r in rows:
        keys = {str(k).strip().lower() for k in r}
        if keys & {"company name", "company", "employer", "organization"} and keys & {"title", "position", "job title", "role"}:
            start, end = _k(r, "started on", "start date", "start", "from"), _k(r, "finished on", "end date", "end", "to")
            out["work"].append({"job_title": _k(r, "title", "position", "job title", "role"), "organization": _k(r, "company name", "company", "employer", "organization"),
                                "location": _k(r, "location") or None, "start_date": start, "end_date": end, "is_current": not end or _is_present(end),
                                "description": _k(r, "description", "summary") or None})
        elif keys & {"school name", "school", "institution", "university"}:
            out["education"].append({"degree": _k(r, "degree name", "degree", "qualification") or "Studies", "school": _k(r, "school name", "school", "institution", "university"),
                                     "field_of_study": _k(r, "field of study", "notes", "course") or None,
                                     "start_date": _k(r, "start date", "started on", "start"), "end_date": _k(r, "end date", "finished on", "end")})
        elif keys & {"authority", "issuer", "issuing organization"} or (keys & {"license number", "credential id"}):
            out["certifications"].append({"name": _k(r, "name", "certification", "title"), "issuer": _k(r, "authority", "issuer", "issuing organization") or None,
                                          "issued_at": _k(r, "started on", "issued", "issue date", "date") or None,
                                          "credential_id": _k(r, "license number", "credential id") or None, "credential_url": _k(r, "url", "credential url") or None})
        elif keys & {"name", "skill", "skills"} and len(keys - {"__sheet"}) <= 3:
            v = _k(r, "name", "skill", "skills")
            out["skills"].extend(s.strip() for s in re.split(r"[;,]", v) if s.strip())
    if not any(out[k] for k in ("work", "education", "certifications", "skills")):
        return None
    out["parser"] = "structured"
    return out


# ─── Heuristic text parser ───────────────────────────────────────────────────

SECTION = {
    "summary": r"(professional\s+)?(summary|profile|about( me)?|objective|career objective|personal statement)",
    "work": r"(work|professional|employment|career)?\s*(experience|history)|employment|work history|positions?|experience",
    "education": r"education(al background)?|academic( background| qualifications)?|qualifications",
    "skills": r"(key |core |technical )?(skills|competenc(ies|e)|expertise|tools|technologies)( & tools)?",
    "certifications": r"certifications?|licen[cs]es?( & certifications)?|certificates|professional (development|certifications)|courses",
    "other": r"references|referees|hobbies|interests|languages|projects|awards|publications|volunteer(ing)?",
}
DEGREE_RE = re.compile(r"(?i)\b(ph\.?\s?d|doctorate|m\.?\s?sc|msc|m\.?\s?a\b|mba|m\.?\s?eng|b\.?\s?sc|bsc|b\.?\s?a\b|b\.?\s?eng|beng|b\.?\s?tech|btech|llb|bachelor'?s?|master'?s?|hnd|ond|nce|diploma|associate degree|waec|ssce|neco|a-?levels?)\b")
SCHOOL_RE = re.compile(r"(?i)\b(university|polytechnic|college|institute|school|academy|unilag|unn|oau|ui\b|abu\b)\b")
CERT_HINT = re.compile(r"(?i)\b(certified|certification|certificate|pmp|acca|ican|cipm|cfa|ccna|comptia|aws certified|scrum master|psm|cism|cissp|prince2)\b")
BULLET = re.compile(r"^\s*(?:[-•*▪●◦·‣–➢➤✓✔○■□►\uf000-\uf8ff]\s*|\d{1,2}[.)]\s+)(?=\S)")
SPLIT_TITLE = re.compile(r"\s+(?:at|@)\s+|\s*[|•·—–]\s*|\s+-\s+|,\s+")


def _section_of(line: str) -> str | None:
    t = re.sub(r"[^a-z& ]", "", line.strip().lower()).strip()
    if not t or len(t) > 40:
        return None
    for k, pat in SECTION.items():
        if re.fullmatch(pat, t):
            return k
    return None


def _clean(s: str) -> str:
    return re.sub(r"\s+", " ", BULLET.sub("", s or "")).strip(" -–—|•·,")


def _empty() -> dict:
    return {"headline": None, "summary": None, "work": [], "education": [], "skills": [], "certifications": []}


def heuristic(text: str) -> dict:
    out = _empty()
    lines = [l.rstrip() for l in text.replace("\r", "").split("\n")]
    sections: dict[str, list[str]] = {"top": []}
    cur = "top"
    for l in lines:
        s = _section_of(l)
        if s:
            cur = s
            sections.setdefault(cur, [])
            continue
        sections.setdefault(cur, []).append(l)
    top = [l.strip() for l in sections.get("top", []) if l.strip()]
    # headline: a short non-contact line near the top
    for l in top[1:6]:
        if not re.search(r"@|\+?\d[\d\s()-]{7,}|https?://|linkedin", l, re.I) and 1 <= len(l.split()) <= 12 and not RANGE_RE.search(l):
            out["headline"] = _clean(l)[:160]
            break
    if sections.get("summary"):
        out["summary"] = " ".join(_clean(l) for l in sections["summary"] if l.strip())[:900] or None

    # work: anchor on date ranges
    wl = sections.get("work") or (lines if not sections.get("education") else [])
    i = 0
    cur_job: dict | None = None
    while i < len(wl):
        l = wl[i].strip()
        m = RANGE_RE.search(l)
        if m:
            head = _clean(RANGE_RE.sub("", l))
            prev = [_clean(x) for x in wl[max(0, i - 2):i] if x.strip() and not BULLET.match(x)]
            if cur_job and prev and cur_job.get("_bullets") and prev[-1] in cur_job["_bullets"]:
                prev = prev[-1:]
            cand = head or (prev[-1] if prev else "")
            parts = [p for p in SPLIT_TITLE.split(cand) if p] if cand else []
            title, org = (parts[0], parts[1]) if len(parts) >= 2 else (cand, "")
            if not org and prev:
                org_line = prev[-2] if head == "" and len(prev) >= 2 else (prev[-1] if head else "")
                if org_line and org_line != title:
                    title, org = (org_line, title) if head == "" else (title, org_line)
            cur_job = {"job_title": title[:200], "organization": org[:200], "start_date": m.group(1), "end_date": m.group(2),
                       "is_current": _is_present(m.group(2)), "responsibilities": [], "achievements": [], "_bullets": []}
            out["work"].append(cur_job)
        elif cur_job and l:
            c = _clean(l)
            if c and (BULLET.match(wl[i]) or len(c.split()) > 6):
                (cur_job["achievements"] if re.search(r"\d|%|₦|\$", c) else cur_job["responsibilities"]).append(c[:300])
                cur_job["_bullets"].append(c)
        i += 1
    for j in out["work"]:
        j.pop("_bullets", None)
        j["responsibilities"] = j["responsibilities"][:8]
        j["achievements"] = j["achievements"][:8]

    # education
    el = sections.get("education") or []
    for idx, l in enumerate(el):
        c = _clean(l)
        if not c:
            continue
        if DEGREE_RE.search(c) or SCHOOL_RE.search(c):
            yrs = re.findall(r"\b(19[5-9]\d|20\d{2})\b", c + " " + " ".join(el[idx + 1:idx + 2]))
            rng = RANGE_RE.search(c)
            body = _clean(RANGE_RE.sub("", c))
            parts = [p for p in re.split(r"\s*[|•·—–,]\s*|\s+-\s+", body) if p]
            degree = next((p for p in parts if DEGREE_RE.search(p)), None)
            school = next((p for p in parts if SCHOOL_RE.search(p) and p != degree), None)
            if not school and idx + 1 < len(el) and SCHOOL_RE.search(el[idx + 1]):
                school = _clean(el[idx + 1])
            if not degree and not school:
                continue
            if out["education"] and not degree and out["education"][-1]["school"] in ("", None):
                out["education"][-1]["school"] = school
                continue
            if degree is None and out["education"] and out["education"][-1].get("_need_school") is False:
                pass
            field = None
            if degree:
                fm = re.search(r"(?i)\b(?:in|of)\s+([A-Za-z &]+)$", degree)
                if not fm:
                    dm = DEGREE_RE.search(degree)
                    rest = degree[dm.end():].strip(" .,-") if dm and dm.start() == 0 else ""
                    if 2 < len(rest) < 80:
                        field, degree = rest, degree[:dm.end()].strip()
                else:
                    field, degree = fm.group(1).strip(), degree[:fm.start()].strip()
            out["education"].append({"degree": (degree or "Studies")[:200], "school": (school or "")[:200], "field_of_study": field,
                                     "start_date": rng.group(1) if rng else (yrs[0] if len(yrs) > 1 else None),
                                     "end_date": rng.group(2) if rng else (yrs[-1] if yrs else None)})
    merged: list[dict] = []
    for e in out["education"]:  # "School" line followed by "Degree, years" line (or the reverse) → one entry
        p = merged[-1] if merged else None
        if p and ((p["degree"] == "Studies" and not e["school"]) or (not p["school"] and e["degree"] == "Studies")):
            p.update({"degree": p["degree"] if p["degree"] != "Studies" else e["degree"], "school": p["school"] or e["school"],
                      "field_of_study": p["field_of_study"] or e["field_of_study"],
                      "start_date": p["start_date"] or e["start_date"], "end_date": p["end_date"] or e["end_date"]})
            continue
        merged.append(e)
    out["education"] = [e for e in merged if e["school"] or e["degree"] != "Studies"][:6]

    # skills
    sk: list[str] = []
    for l in sections.get("skills") or []:
        c = re.sub(r"^[A-Za-z /&]{2,30}:\s*", "", _clean(l))
        sk.extend(p.strip() for p in re.split(r"[,;|•·]|\s{2,}|\s/\s", c) if 1 < len(p.strip()) <= 40)
    sk.extend(vocab.find_terms(text[:20000]))
    out["skills"] = sk

    # certifications
    cl = sections.get("certifications") or []
    pool = cl if cl else [l for l in lines if CERT_HINT.search(l) and len(l) < 140]
    for l in pool:
        c = _clean(l)
        if not c or len(c) < 3:
            continue
        yr = re.search(r"\b(19[5-9]\d|20\d{2})\b", c)
        body = _clean(re.sub(r"\(?\b(19[5-9]\d|20\d{2})\b\)?", "", c))
        parts = [p for p in re.split(r"\s*[|•·,]\s*", body) if p]
        if len(parts) < 2:
            parts = [p for p in re.split(r"\s+[—–-]\s+", body) if p]
        out["certifications"].append({"name": parts[0][:200], "issuer": parts[1][:200] if len(parts) > 1 else None,
                                      "issued_at": yr.group(1) if yr else None})
    out["certifications"] = out["certifications"][:12]
    out["parser"] = "heuristic"
    return out


# ─── AI parser (OpenAI-compatible, cheap) ────────────────────────────────────

AI_SYSTEM = """Extract career facts from the document. Only include facts present in the text — never invent.
Return JSON exactly:
{"headline": str|null, "summary": str|null,
 "work": [{"job_title": str, "organization": str, "location": str|null, "start_date": "YYYY-MM"|null, "end_date": "YYYY-MM"|null,
           "is_current": bool, "description": str|null, "responsibilities": [str], "achievements": [str]}],
 "education": [{"degree": str, "school": str, "field_of_study": str|null, "start_date": "YYYY-MM"|null, "end_date": "YYYY-MM"|null, "grade": str|null}],
 "skills": [str],
 "certifications": [{"name": str, "issuer": str|null, "issued_at": "YYYY-MM"|null, "credential_id": str|null}]}
Achievements = quantified results; responsibilities = duties. Keep each bullet under 200 chars. Max 12 roles, 40 skills."""


async def ai_parse(text: str) -> dict | None:
    if not llm.llm_available():
        return None
    try:
        res = await llm.chat_json(AI_SYSTEM, text[:14000], temperature=0.1, max_tokens=2200)
    except Exception as e:  # noqa: BLE001
        log.info("AI import parse failed, using heuristics: %s", e)
        return None
    if not isinstance(res, dict) or not any(res.get(k) for k in ("work", "education", "skills", "certifications")):
        return None
    out = _empty()
    out.update({k: res.get(k) for k in out if res.get(k)})
    out["parser"] = "ai"
    return out


# ─── Normalise + dedupe against the profile ──────────────────────────────────


def _n(s: Any) -> str:
    return vocab.norm(str(s or ""))


def _same_org(a: str, b: str) -> bool:
    a, b = _n(a), _n(b)
    if not a or not b:
        return False
    strip = lambda x: re.sub(r"\b(ltd|limited|plc|inc|llc|nigeria|ng|group|company|co)\b", "", x).strip()
    a, b = strip(a), strip(b)
    return a == b or (len(a) > 3 and (a in b or b in a))


def _same_title(a: str, b: str) -> bool:
    a, b = set(_n(a).split()), set(_n(b).split())
    return bool(a and b) and len(a & b) / max(1, min(len(a), len(b))) >= 0.6


def _list(v: Any) -> list[str]:
    if not v:
        return []
    if isinstance(v, str):
        v = [v]
    return [str(x).strip()[:300] for x in v if str(x).strip()]


def build_proposal(profile: Profile, raw: dict) -> dict:
    today = date.today()
    prop: dict[str, Any] = {"headline": None, "summary": None, "work": [], "education": [], "skills": [], "certifications": []}
    existing_work = [e for e in profile.timeline_entries if (e.entry_type or "work") == "work"]
    existing_edu = [e for e in profile.timeline_entries if e.entry_type == "education"]

    def _entry_common(it: dict, title: str, org: str) -> dict:
        start, end = parse_date(it.get("start_date")), parse_date(it.get("end_date"))
        current = bool(it.get("is_current")) or (_is_present(it.get("end_date")) if it.get("end_date") else False)
        guessed = False
        if not start:
            start, guessed = (end or today.replace(month=1, day=1)), True
        if end and end < start:
            start, end = end, start
        return {"job_title": title[:255], "organization": org[:255], "start_date": start.isoformat(),
                "end_date": None if current or not end else end.isoformat(), "is_current": current, "date_guessed": guessed}

    seen: set[tuple] = set()
    for i, it in enumerate(raw.get("work") or []):
        title, org = str(it.get("job_title") or "").strip(), str(it.get("organization") or "").strip()
        if not title or len(title) > 200:
            continue
        row = _entry_common(it, title, org or "—")
        key = (_n(title), _n(org), row["start_date"][:4])
        if key in seen:
            continue
        seen.add(key)
        row.update({"key": f"w{i}", "location": it.get("location") or None, "description": (it.get("description") or None),
                    "responsibilities": _list(it.get("responsibilities"))[:8], "achievements": _list(it.get("achievements"))[:8]})
        match = next((e for e in existing_work if _same_org(e.organization, org) and (_same_title(e.job_title, title) or e.start_date.year == int(row["start_date"][:4]))), None)
        if match:
            changes = []
            if not match.description and row["description"]:
                changes.append("description")
            new_ach = [a for a in row["achievements"] if _n(a) not in {_n(x) for x in (match.achievements or [])}]
            new_resp = [a for a in row["responsibilities"] if _n(a) not in {_n(x) for x in (match.responsibilities or [])}]
            if new_ach:
                changes.append(f"{len(new_ach)} achievement{'s' if len(new_ach) != 1 else ''}")
            if new_resp:
                changes.append(f"{len(new_resp)} responsibilit{'ies' if len(new_resp) != 1 else 'y'}")
            if not match.location and row["location"]:
                changes.append("location")
            row.update({"status": "update" if changes else "duplicate", "match_id": match.id, "changes": changes,
                        "achievements": new_ach, "responsibilities": new_resp})
        else:
            row.update({"status": "new", "match_id": None, "changes": []})
        prop["work"].append(row)

    for i, it in enumerate(raw.get("education") or []):
        degree, school = str(it.get("degree") or "").strip(), str(it.get("school") or "").strip()
        if not (degree or school):
            continue
        row = _entry_common(it, degree or "Studies", school or "—")
        row["is_current"] = bool(it.get("is_current")) and not it.get("end_date")
        if not row["is_current"] and not row["end_date"] and it.get("end_date"):
            row["end_date"] = (parse_date(it.get("end_date")) or today).isoformat()
        row.update({"key": f"e{i}", "field_of_study": it.get("field_of_study") or None, "grade": it.get("grade") or None})
        match = next((e for e in existing_edu if _same_org(e.organization, school) or (_same_title(e.job_title, degree) and _n(school)[:6] == _n(e.organization)[:6])), None)
        if match:
            changes = [label for f, label in (("field_of_study", "field of study"), ("grade", "grade")) if row.get(f) and not getattr(match, f, None)]
            row.update({"status": "update" if changes else "duplicate", "match_id": match.id, "changes": changes})
        else:
            row.update({"status": "new", "match_id": None, "changes": []})
        prop["education"].append(row)

    have = {vocab.canon(s.name) for s in profile.skills}
    added: set[str] = set()
    for s in raw.get("skills") or []:
        name = str(s if not isinstance(s, dict) else s.get("name") or "").strip().strip(".")
        if not name or len(name) > 60 or len(name.split()) > 5:
            continue
        c = vocab.canon(name)
        if c in added or not c:
            continue
        added.add(c)
        from app.services.insights import pretty
        display = name if any(ch.isupper() for ch in name) else pretty(name)
        prop["skills"].append({"name": display[:120], "status": "duplicate" if c in have else "new"})
    prop["skills"] = sorted(prop["skills"], key=lambda x: x["status"] != "new")[:50]

    have_c = [(_n(c.name), c) for c in profile.certifications]
    for i, it in enumerate(raw.get("certifications") or []):
        name = str(it.get("name") or "").strip()
        if not name or len(name) > 200:
            continue
        nn = _n(name)
        dup = next((c for n, c in have_c if n == nn or (len(nn) > 6 and (nn in n or n in nn))), None)
        issued = parse_date(it.get("issued_at"))
        prop["certifications"].append({"key": f"c{i}", "name": name[:255], "issuer": (it.get("issuer") or None), "issued_at": issued.isoformat() if issued else None,
                                       "credential_id": it.get("credential_id") or None, "credential_url": it.get("credential_url") or None,
                                       "status": "duplicate" if dup else "new", "match_id": dup.id if dup else None})

    h, s = (raw.get("headline") or "").strip(), (raw.get("summary") or "").strip()
    if h and _n(h) != _n(profile.headline):
        prop["headline"] = {"value": h[:255], "current": profile.headline, "replace": bool(profile.headline)}
    if s and len(s) > 30 and _n(s) != _n(profile.summary):
        prop["summary"] = {"value": s[:2000], "current": profile.summary, "replace": bool(profile.summary)}
    prop["counts"] = {k: sum(1 for x in prop[k] if x["status"] != "duplicate") for k in ("work", "education", "skills", "certifications")}
    return prop


# ─── Apply ───────────────────────────────────────────────────────────────────


def _org_id(db: Session, name: str, kind: str) -> tuple[int | None, str]:
    from app.services.orgs import find_by_name
    try:
        org = find_by_name(db, name or "", kind) if name and name != "—" else None
    except Exception:  # noqa: BLE001
        org = None
    return (org.id, org.name) if org else (None, name)


def apply(db: Session, profile: Profile, body: dict) -> dict:
    """Create/merge the approved items. No verification requests are sent — users request those explicitly."""
    done = {"work": 0, "education": 0, "skills": 0, "certifications": 0, "updated": 0, "headline": False, "summary": False}
    entries = {e.id: e for e in profile.timeline_entries}
    for kind in ("work", "education"):
        for it in body.get(kind) or []:
            try:
                start = parse_date(it.get("start_date")) or date.today()
                end = None if it.get("is_current") else parse_date(it.get("end_date"))
            except Exception:  # noqa: BLE001
                continue
            title, org = str(it.get("job_title") or "").strip()[:255], str(it.get("organization") or "").strip()[:255] or "—"
            if not title:
                continue
            m = entries.get(it.get("match_id")) if it.get("match_id") else None
            if m and m.profile_id == profile.id:
                if kind == "work":
                    if it.get("description") and not m.description:
                        m.description = it["description"]
                    if it.get("location") and not m.location:
                        m.location = it["location"]
                    for f in ("achievements", "responsibilities"):
                        cur = list(getattr(m, f) or [])
                        seen = {_n(x) for x in cur}
                        cur += [x for x in _list(it.get(f)) if _n(x) not in seen]
                        setattr(m, f, cur[:20])
                else:
                    for f in ("field_of_study", "grade"):
                        if it.get(f) and not getattr(m, f, None):
                            setattr(m, f, str(it[f])[:255])
                done["updated"] += 1
                continue
            oid, oname = _org_id(db, org, "school" if kind == "education" else "company")
            e = TimelineEntry(profile_id=profile.id, entry_type=kind, job_title=title, organization=oname[:255], organization_id=oid,
                              start_date=start, end_date=end, is_current=bool(it.get("is_current")),
                              location=((it.get("location") or "")[:120] or None), description=(it.get("description") or None),
                              responsibilities=_list(it.get("responsibilities"))[:20] or None, achievements=_list(it.get("achievements"))[:20] or None,
                              field_of_study=(it.get("field_of_study") or None), grade=(str(it["grade"])[:64] if it.get("grade") else None))
            db.add(e)
            done[kind] += 1
    have = {vocab.canon(s.name) for s in profile.skills}
    for s in body.get("skills") or []:
        name = str(s.get("name") if isinstance(s, dict) else s or "").strip()[:120]
        if name and vocab.canon(name) not in have:
            have.add(vocab.canon(name))
            db.add(Skill(profile_id=profile.id, name=name, category="technical"))
            done["skills"] += 1
    have_c = {_n(c.name) for c in profile.certifications}
    for c in body.get("certifications") or []:
        name = str(c.get("name") or "").strip()[:255]
        if not name or _n(name) in have_c:
            continue
        have_c.add(_n(name))
        oid, oname = _org_id(db, c.get("issuer") or "", "issuer")
        db.add(Certification(profile_id=profile.id, name=name, issuer=(oname or None) if c.get("issuer") else None, issuer_org_id=oid,
                             issued_at=parse_date(c.get("issued_at")), credential_id=(c.get("credential_id") or None),
                             credential_url=(c.get("credential_url") or None)))
        done["certifications"] += 1
    if isinstance(body.get("headline"), str) and body["headline"].strip():
        profile.headline = body["headline"].strip()[:255]
        done["headline"] = True
    if isinstance(body.get("summary"), str) and body["summary"].strip():
        profile.summary = body["summary"].strip()[:4000]
        done["summary"] = True
    db.commit()
    db.refresh(profile)
    return done
